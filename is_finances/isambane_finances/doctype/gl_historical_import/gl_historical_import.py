# Copyright (c) 2026, Isambane Mining and contributors
# For license information, please see license.txt

import calendar
import json
import os
import re
import tempfile
import zipfile
from collections import defaultdict
from datetime import date
from io import BytesIO

import frappe
from frappe.model.document import Document
from frappe.utils import cstr, flt, get_files_path, get_site_path


MONTH_ALIASES = {
    "jan": 1,
    "january": 1,
    "feb": 2,
    "february": 2,
    "mar": 3,
    "march": 3,
    "apr": 4,
    "april": 4,
    "may": 5,
    "jun": 6,
    "june": 6,
    "jul": 7,
    "july": 7,
    "aug": 8,
    "august": 8,
    "sep": 9,
    "sept": 9,
    "september": 9,
    "oct": 10,
    "october": 10,
    "nov": 11,
    "november": 11,
    "dec": 12,
    "december": 12,
}

CONTROL_LABELS = {
    "total revenue": "revenue",
    "total other income": "other_income",
    "total other expense": "expense",
    "profit before tax": "profit_loss",
    "loss before tax": "profit_loss",
}

SECTION_LABELS = {
    "revenue",
    "gross profit",
    "other income",
    "other expense",
    "income statement",
    "statement of comprehensive income year",
    "statement of comprehensive income",
}


class GLHistoricalImport(Document):
    def validate(self):
        self.populate_root_types()

    def populate_root_types(self):
        for row in self.get("gl_historical_lines") or []:
            if row.account_number:
                row.root_type = frappe.db.get_value("Account", row.account_number, "root_type")


@frappe.whitelist()
def preview_gl_historical_import(docname, selected_months=None):
    doc = frappe.get_doc("GL Historical Import", docname)
    selected = _normalise_selected_months(selected_months, doc.import_year)
    parsed = _parse_excel_for_doc(doc, selected, import_rows=False)
    if parsed.get("control_difference_errors"):
        parsed["errors"] = list(parsed.get("errors") or []) + list(parsed.get("control_difference_errors") or [])
    duplicates = _find_existing_company_month_imports(doc, selected)
    if duplicates:
        parsed["warnings"].append(_duplicate_warning_text(duplicates))
    html = _build_summary_html(parsed, title="Preview - GL Historical Import")
    doc.db_set("status", "Previewed", update_modified=True)
    doc.db_set("import_months", ", ".join([m["label"] for m in selected]), update_modified=True)
    return {
        "html": html,
        "selected_months": [m["label"] for m in selected],
        "row_count": parsed["row_count"],
        "errors": parsed["errors"],
        "warnings": parsed["warnings"],
        "duplicates": duplicates,
        "has_duplicates": bool(duplicates),
    }


@frappe.whitelist()
def import_gl_historical(docname, selected_months=None, duplicate_action=None):
    doc = frappe.get_doc("GL Historical Import", docname)
    selected = _normalise_selected_months(selected_months, doc.import_year)

    duplicates = _find_existing_company_month_imports(doc, selected)
    if duplicates and duplicate_action not in ("replace", "keep"):
        parsed = _parse_excel_for_doc(doc, selected, import_rows=False)
        parsed["warnings"].append(_duplicate_warning_text(duplicates))
        return {
            "requires_duplicate_decision": True,
            "duplicates": duplicates,
            "html": _build_summary_html(parsed, title="Duplicate Month Data Found"),
            "errors": [],
            "warnings": parsed["warnings"],
            "imported": 0,
        }

    if duplicates and duplicate_action == "keep":
        parsed = _parse_excel_for_doc(doc, selected, import_rows=False)
        parsed["warnings"].append("Import stopped. Existing company/month data was kept unchanged.")
        return {
            "kept_existing": True,
            "duplicates": duplicates,
            "html": _build_summary_html(parsed, title="Import Stopped - Existing Data Kept"),
            "errors": [],
            "warnings": parsed["warnings"],
            "imported": 0,
        }

    if duplicates and duplicate_action == "replace":
        _replace_existing_company_month_data(doc, selected, duplicates)

    parsed = _parse_excel_for_doc(doc, selected, import_rows=True)

    if parsed["errors"]:
        doc.db_set("status", "Failed", update_modified=True)
        return {
            "html": _build_summary_html(parsed, title="Import Failed - GL Historical Import"),
            "errors": parsed["errors"],
            "warnings": parsed["warnings"],
            "imported": 0,
        }

    if not parsed.get("lines"):
        parsed["errors"].append(
            "No importable account rows were found. Check that the selected month exists in the Excel file and that the Account/Actual columns are present."
        )
        doc.db_set("status", "Failed", update_modified=True)
        return {
            "html": _build_summary_html(parsed, title="Import Failed - No Lines Found"),
            "errors": parsed["errors"],
            "warnings": parsed["warnings"],
            "imported": 0,
        }

    doc.set("gl_historical_lines", [])
    for line in parsed["lines"]:
        doc.append("gl_historical_lines", line)

    doc.import_months = ", ".join([m["label"] for m in selected])
    doc.status = "Imported"
    doc.save(ignore_permissions=True)

    parsed_after_save = _summary_from_imported_doc(doc)
    parsed_after_save["excel_controls"] = parsed["excel_controls"]
    parsed_after_save["warnings"] = parsed["warnings"]
    parsed_after_save["errors"] = _validation_differences(parsed_after_save)

    if parsed_after_save["errors"]:
        doc.db_set("status", "Failed", update_modified=True)
        parsed_after_save["status"] = "Failed"
    else:
        doc.db_set("status", "Imported", update_modified=True)
        parsed_after_save["status"] = "Imported"

    html = _build_summary_html(parsed_after_save, title="Imported - GL Historical Import")
    return {
        "html": html,
        "errors": parsed_after_save["errors"],
        "warnings": parsed_after_save["warnings"],
        "imported": len(doc.get("gl_historical_lines") or []),
    }


@frappe.whitelist()
def check_existing_company_month_data(docname, selected_months=None):
    doc = frappe.get_doc("GL Historical Import", docname)
    selected = _normalise_selected_months(selected_months, doc.import_year)
    duplicates = _find_existing_company_month_imports(doc, selected)
    return {
        "has_duplicates": bool(duplicates),
        "duplicates": duplicates,
        "message": _duplicate_warning_text(duplicates) if duplicates else "",
    }


@frappe.whitelist()
def get_gl_historical_summary(docname):
    doc = frappe.get_doc("GL Historical Import", docname)
    selected = _months_from_doc(doc)

    if not selected or not doc.excel_file:
        parsed = _summary_from_imported_doc(doc)
        return {"html": _build_summary_html(parsed, title="GL Historical Import")}

    try:
        parsed_excel = _parse_excel_for_doc(doc, selected, import_rows=False)
        parsed = _summary_from_imported_doc(doc)
        parsed["excel_controls"] = parsed_excel["excel_controls"]
        parsed["warnings"].extend(parsed_excel.get("warnings") or [])
        parsed["errors"].extend(_validation_differences(parsed))
        return {"html": _build_summary_html(parsed, title="GL Historical Import")}
    except Exception as exc:
        parsed = _summary_from_imported_doc(doc)
        parsed["warnings"].append(f"Could not refresh Excel control totals: {frappe.bold(cstr(exc))}")
        return {"html": _build_summary_html(parsed, title="GL Historical Import")}


def _normalise_selected_months(selected_months, import_year):
    if not import_year:
        frappe.throw("Please enter Import Year before preview/import.")

    if isinstance(selected_months, str):
        try:
            selected_months = json.loads(selected_months)
        except Exception:
            selected_months = [x.strip() for x in selected_months.split(",") if x.strip()]

    selected_months = selected_months or []
    out = []
    seen = set()

    for item in selected_months:
        month_no = None
        if isinstance(item, dict):
            month_no = item.get("month") or item.get("month_no") or item.get("value")
        else:
            text = cstr(item).strip()
            if text.isdigit():
                month_no = int(text)
            else:
                parsed = _parse_month_year(text)
                if parsed:
                    month_no = parsed[0]

        month_no = int(month_no or 0)
        if month_no < 1 or month_no > 12:
            continue
        if month_no in seen:
            continue
        seen.add(month_no)
        out.append({
            "month": month_no,
            "year": int(import_year),
            "label": f"{calendar.month_abbr[month_no]} {int(import_year)}",
            "month_end_date": _month_end_date(int(import_year), month_no),
        })

    if not out:
        frappe.throw("Please select at least one month.")

    return out


def _months_from_doc(doc):
    if not doc.import_months:
        return []
    return _normalise_selected_months(doc.import_months, doc.import_year)


def _parse_excel_for_doc(doc, selected_months, import_rows=False):
    if not doc.company:
        frappe.throw("Please select Company before preview/import.")
    if not doc.excel_file:
        frappe.throw("Please attach the Excel file before preview/import.")

    path = _get_file_path(doc.excel_file)
    wb = _load_workbook(path)
    ws = wb.active
    rows = [[cell.value for cell in row] for row in ws.iter_rows()]

    month_columns = _find_month_columns(rows)
    single_month = _find_single_month_report(rows)
    actual_col = _find_single_actual_column(rows)

    selected_cols = []
    warnings = []
    errors = []

    for selected in selected_months:
        key = (selected["month"], selected["year"])
        col = month_columns.get(key)
        if col is None and single_month and key == single_month and actual_col is not None:
            col = actual_col
        if col is None:
            if single_month:
                single_label = f"{calendar.month_abbr[single_month[0]]} {single_month[1]}"
                errors.append(
                    f"Could not find an Excel column for {selected['label']}. "
                    f"This appears to be a single-month Excel report for {single_label}. "
                    f"Please select only {single_label}, or attach a multi-month report."
                )
            else:
                available = ", ".join(
                    f"{calendar.month_abbr[m]} {y}" for (m, y) in sorted(month_columns.keys(), key=lambda x: (x[1], x[0]))
                ) or "none"
                errors.append(
                    f"Could not find an Excel column for {selected['label']}. Available month columns found: {available}."
                )
        else:
            selected_cols.append((selected, col))

    if single_month:
        warnings.append(f"Detected single-month report: {calendar.month_abbr[single_month[0]]} {single_month[1]}.")
    elif month_columns:
        detected = ", ".join(
            f"{calendar.month_abbr[m]} {y}" for (m, y) in sorted(month_columns.keys(), key=lambda x: (x[1], x[0]))
        )
        warnings.append(f"Detected multi-month report columns: {detected}.")

    excel_controls = _extract_excel_control_totals(rows, selected_cols)
    lines = []
    account_cache = {}
    cost_center_cache = {}
    root_type_cache = {}

    for row_idx, row in enumerate(rows, start=1):
        label = cstr(row[0] if row else "").strip()
        parsed_account = _parse_account_label(label)
        if not parsed_account:
            continue

        account_code = parsed_account["account_code"]
        cost_center_code = parsed_account["cost_center_code"]
        cost_center_was_explicit = parsed_account.get("cost_center_was_explicit")

        account_name = account_cache.get(account_code)
        if account_name is None:
            account_name = _resolve_account(account_code, doc.company)
            account_cache[account_code] = account_name

        if not account_name:
            errors.append(f"Row {row_idx}: Account {account_code} could not be found for company {doc.company}.")
            continue

        root_type = root_type_cache.get(account_name)
        if root_type is None:
            root_type = frappe.db.get_value("Account", account_name, "root_type")
            root_type_cache[account_name] = root_type

        if root_type not in ("Income", "Expense"):
            errors.append(
                f"Row {row_idx}: Account {account_name} has root_type {root_type or 'blank'}, "
                "but only Income and Expense accounts may be imported."
            )
            continue

        cost_center_name = _get_resolved_cost_center_for_row(
            cost_center_code=cost_center_code,
            company=doc.company,
            row_idx=row_idx,
            warnings=warnings,
            errors=errors,
            cache=cost_center_cache,
            cost_center_was_explicit=cost_center_was_explicit,
        )
        if not cost_center_name:
            continue

        for selected, col in selected_cols:
            actual = _cell_amount(row, col)
            if actual == 0:
                continue
            lines.append({
                "month_end_date": selected["month_end_date"],
                "account_number": account_name,
                "cost_center_number": cost_center_name,
                "actual": actual,
                "root_type": root_type,
                "excel_month_label": selected["label"],
                "source_row": row_idx,
            })

    imported = _calculate_imported_totals(lines)

    parsed = {
        "status": "Previewed",
        "selected_months": selected_months,
        "excel_controls": excel_controls,
        "imported_totals": imported,
        "lines": lines if import_rows else [],
        "row_count": len(lines),
        # errors here are fatal parse/master-data errors only.
        # Control-total differences are calculated later and must not prevent
        # rows from being saved, because the user must be able to review the
        # imported lines against the Excel controls.
        "errors": errors,
        "warnings": warnings,
    }

    difference_errors = _validation_differences(parsed)
    if difference_errors:
        parsed["control_difference_errors"] = difference_errors
        parsed["warnings"].append(
            "Control totals do not yet balance. Preview/import will show the differences after line extraction."
        )
    return parsed


def _find_existing_company_month_imports(doc, selected_months):
    """Return other GL Historical Import records that already contain lines for company/month."""
    if not doc.company or not selected_months:
        return []

    month_end_dates = [m["month_end_date"] for m in selected_months]
    if not month_end_dates:
        return []

    selected_by_date = {m["month_end_date"]: m["label"] for m in selected_months}
    parent_dt = "tabGL Historical Import"
    child_dt = "tabGL Historical Line"

    rows = frappe.db.sql(
        f"""
        select
            child.parent as import_name,
            parent.status as status,
            parent.import_year as import_year,
            parent.import_months as import_months,
            child.month_end_date as month_end_date,
            count(*) as line_count,
            sum(child.actual) as amount_total
        from `{child_dt}` child
        inner join `{parent_dt}` parent on parent.name = child.parent
        where child.parenttype = 'GL Historical Import'
          and child.parent != %(current_doc)s
          and parent.company = %(company)s
          and child.month_end_date in %(month_end_dates)s
        group by child.parent, parent.status, parent.import_year, parent.import_months, child.month_end_date
        order by child.month_end_date, child.parent
        """,
        {
            "current_doc": doc.name,
            "company": doc.company,
            "month_end_dates": tuple(month_end_dates),
        },
        as_dict=True,
    )

    duplicates = []
    for row in rows:
        month_end_date = cstr(row.month_end_date)
        duplicates.append({
            "import_name": row.import_name,
            "status": row.status,
            "month_end_date": month_end_date,
            "month_label": selected_by_date.get(month_end_date, month_end_date),
            "line_count": int(row.line_count or 0),
            "amount_total": flt(row.amount_total, 2),
        })
    return duplicates


def _duplicate_warning_text(duplicates):
    if not duplicates:
        return ""
    parts = []
    for item in duplicates:
        parts.append(
            f"{item.get('month_label')} already exists in {item.get('import_name')} "
            f"({item.get('line_count')} lines)."
        )
    return "Existing company/month data found: " + " ".join(parts)


def _replace_existing_company_month_data(doc, selected_months, duplicates=None):
    """Delete matching month lines from other import batches for the same company."""
    month_end_dates = [m["month_end_date"] for m in selected_months]
    if not month_end_dates:
        return

    if duplicates is None:
        duplicates = _find_existing_company_month_imports(doc, selected_months)

    affected_parents = sorted(set(d.get("import_name") for d in duplicates if d.get("import_name")))
    if not affected_parents:
        return

    frappe.db.sql(
        """
        delete child
        from `tabGL Historical Line` child
        inner join `tabGL Historical Import` parent on parent.name = child.parent
        where child.parenttype = 'GL Historical Import'
          and child.parent != %(current_doc)s
          and parent.company = %(company)s
          and child.month_end_date in %(month_end_dates)s
        """,
        {
            "current_doc": doc.name,
            "company": doc.company,
            "month_end_dates": tuple(month_end_dates),
        },
    )

    for parent_name in affected_parents:
        _refresh_parent_import_months(parent_name)


def _refresh_parent_import_months(parent_name):
    rows = frappe.get_all(
        "GL Historical Line",
        filters={"parent": parent_name, "parenttype": "GL Historical Import"},
        fields=["distinct month_end_date"],
        order_by="month_end_date asc",
    )
    labels = []
    for row in rows:
        month_end = row.get("month_end_date")
        if not month_end:
            continue
        if isinstance(month_end, str):
            year = int(month_end[:4])
            month = int(month_end[5:7])
        else:
            year = month_end.year
            month = month_end.month
        labels.append(f"{calendar.month_abbr[month]} {year}")

    frappe.db.set_value(
        "GL Historical Import",
        parent_name,
        "import_months",
        ", ".join(labels),
        update_modified=True,
    )


def _get_file_path(file_url):
    if not file_url:
        frappe.throw("No Excel file attached.")

    file_doc = frappe.db.get_value("File", {"file_url": file_url}, ["is_private", "file_name"], as_dict=True)
    if file_doc:
        if file_doc.is_private:
            return get_site_path("private", "files", file_doc.file_name)
        return get_files_path(file_doc.file_name)

    if file_url.startswith("/private/files/"):
        return get_site_path("private", "files", os.path.basename(file_url))
    if file_url.startswith("/files/"):
        return get_files_path(os.path.basename(file_url))

    frappe.throw(f"Could not resolve attached file path: {file_url}")


def _load_workbook(path):
    """Load Excel workbook, repairing XML attributes that some openpyxl versions reject.

    Some Excel exports contain page/workbook setup attributes such as firstPageNo,
    WindowWidth, WindowHeight, XWindow, and YWindow. openpyxl can raise a TypeError
    before any row data is read. We therefore create a temporary sanitized copy and
    retry when that known incompatibility is detected.
    """
    from openpyxl import load_workbook

    try:
        return load_workbook(path, data_only=True, read_only=True)
    except TypeError as exc:
        if not _is_openpyxl_xml_attribute_error(exc):
            raise
        fixed = _make_openpyxl_safe_copy(path)
        return load_workbook(fixed, data_only=True, read_only=True)


def _is_openpyxl_xml_attribute_error(exc):
    text = cstr(exc)
    known_attributes = (
        "firstPageNo",
        "WindowWidth",
        "WindowHeight",
        "XWindow",
        "YWindow",
    )
    return any(attr in text for attr in known_attributes)


def _make_openpyxl_safe_copy(path):
    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")
    tmp.close()

    workbook_replacements = {
        b"WindowWidth": b"windowWidth",
        b"WindowHeight": b"windowHeight",
        b"XWindow": b"xWindow",
        b"YWindow": b"yWindow",
    }

    with zipfile.ZipFile(path, "r") as zin, zipfile.ZipFile(tmp.name, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)

            if item.filename == "xl/workbook.xml":
                for old, new in workbook_replacements.items():
                    data = data.replace(old, new)

            # Some exports write firstPageNo on worksheet pageSetup. openpyxl 3.x
            # does not accept that keyword in PrintPageSetup.__init__(). Removing
            # only this optional print-setting attribute leaves the actual cell data
            # unchanged.
            if item.filename.startswith("xl/worksheets/") and item.filename.endswith(".xml"):
                data = re.sub(rb'\sfirstPageNo="[^"]*"', b"", data)

            zout.writestr(item, data)

    return tmp.name


def _find_month_columns(rows):
    out = {}
    for row in rows[:30]:
        for idx, value in enumerate(row):
            parsed = _parse_month_year(value)
            if parsed:
                out[parsed] = idx
    return out


def _find_single_month_report(rows):
    for row in rows[:15]:
        text = " ".join([cstr(x) for x in row if x])
        m = re.search(r"For:\s*([A-Za-z]+)\s+(\d{4})", text, flags=re.I)
        if m:
            return _parse_month_year(f"{m.group(1)} {m.group(2)}")
    return None


def _find_single_actual_column(rows):
    for row in rows[:25]:
        if any(cstr(x).strip().lower() == "account" for x in row):
            for idx, value in enumerate(row):
                if cstr(value).strip().lower() == "actual":
                    return idx
    for row in rows[:25]:
        for idx, value in enumerate(row):
            if cstr(value).strip().lower() == "actual":
                return idx
    return None


def _parse_month_year(value):
    text = cstr(value).strip()
    if not text:
        return None
    m = re.match(r"^([A-Za-z]+)\s+(\d{4})$", text)
    if not m:
        return None
    month = MONTH_ALIASES.get(m.group(1).strip().lower())
    if not month:
        return None
    return (month, int(m.group(2)))


def _month_end_date(year, month):
    return date(year, month, calendar.monthrange(year, month)[1]).isoformat()


def _parse_account_label(label):
    """
    Parse account rows from the Statement of Comprehensive Income.

    Expected Excel formats:
        1000>026 (Transport income - WON)
            account_number = 1000
            cost_center_number = WON

        3160>003 (Minor assets - Various - OTH)
            account_number = 3160
            cost_center_number = OTH

        1900 (Interest Received)
            account_number = 1900
            cost_center_number = HO

    Important: the number after > is NOT used as the cost center.
    The cost center is the final code inside brackets after the last hyphen.
    If no final hyphen/code exists, default to HO.
    """
    label = cstr(label).strip()
    if not label:
        return None

    label_lower = label.lower()
    if label_lower in SECTION_LABELS:
        return None
    if label_lower.startswith("total "):
        return None
    if "profit before tax" in label_lower or "loss before tax" in label_lower:
        return None

    m = re.match(r"^(\d+)(?:\s*>\s*([A-Za-z0-9_-]+))?\s*\((.*?)\)\s*$", label)
    if not m:
        return None

    account_code = m.group(1).strip()
    excel_sub_account = (m.group(2) or "").strip() or None
    description = (m.group(3) or "").strip()
    cost_center_code, cost_center_was_explicit = _extract_cost_center_from_description(description)

    return {
        "account_code": account_code,
        "cost_center_code": cost_center_code,
        "cost_center_was_explicit": cost_center_was_explicit,
        "description": description,
        "excel_sub_account": excel_sub_account,
    }


def _extract_cost_center_from_description(description):
    """Return (cost_center_code, was_explicit).

    Business rule confirmed for the Excel Statement of Comprehensive Income:
      - A cost center is only explicit when the bracket description ends with
        a hyphen followed by exactly three capital letters.
      - In the full Excel label this looks like: `... - WON)`.
      - Anything else is description text and must default silently to HO.

    Examples:
        Transport income - WON -> WON
        Minor assets - Various - OTH -> OTH
        Interest Received -> HO
        Bank Charges - BANK -> HO, because BANK is not exactly 3 letters
        Co Contribution - YOUTHWORKERS -> HO, because YOUTHWORKERS is not 3 letters
    """
    text = cstr(description).strip()
    if not text:
        return "HO", False

    match = re.search(r"-\s*([A-Z]{3})\s*$", text)
    if not match:
        return "HO", False

    return match.group(1), True


def _cell_amount(row, col):
    if col is None or col >= len(row):
        return 0
    value = row[col]
    if value in (None, ""):
        return 0
    if isinstance(value, str):
        value = value.replace(",", "").replace("R", "").strip()
        if value in ("-", ""):
            return 0
    return flt(value, 2)


def _extract_excel_control_totals(rows, selected_cols):
    controls = defaultdict(lambda: {"income": 0, "expense": 0, "profit_loss": 0})

    for row in rows:
        label = cstr(row[0] if row else "").strip().lower()
        key = CONTROL_LABELS.get(label)
        if not key:
            continue
        for selected, col in selected_cols:
            amount = _cell_amount(row, col)
            month_label = selected["label"]
            if key in ("revenue", "other_income"):
                controls[month_label]["income"] += amount
            elif key == "expense":
                controls[month_label]["expense"] += amount
            elif key == "profit_loss":
                controls[month_label]["profit_loss"] = amount

    return dict(controls)


def _calculate_imported_totals(lines):
    totals = defaultdict(lambda: {"income": 0, "expense": 0, "profit_loss": 0})
    for line in lines:
        month_label = line.get("excel_month_label")
        actual = flt(line.get("actual"), 2)
        if line.get("root_type") == "Income":
            totals[month_label]["income"] += actual
        elif line.get("root_type") == "Expense":
            totals[month_label]["expense"] += actual

    for month_label, values in totals.items():
        values["profit_loss"] = flt(values["income"] - values["expense"], 2)
    return dict(totals)


def _summary_from_imported_doc(doc):
    lines = []
    for row in doc.get("gl_historical_lines") or []:
        lines.append({
            "excel_month_label": row.excel_month_label,
            "actual": row.actual,
            "root_type": row.root_type,
        })
    return {
        "status": doc.status,
        "selected_months": _months_from_doc(doc),
        "excel_controls": {},
        "imported_totals": _calculate_imported_totals(lines),
        "lines": [],
        "row_count": len(lines),
        "errors": [],
        "warnings": [],
    }


def _validation_differences(parsed):
    errors = []
    controls = parsed.get("excel_controls") or {}
    imported = parsed.get("imported_totals") or {}

    for month_label, excel in controls.items():
        calc = imported.get(month_label, {})
        for key, label in (("income", "Income"), ("expense", "Expense"), ("profit_loss", "Profit/Loss")):
            diff = flt(calc.get(key), 2) - flt(excel.get(key), 2)
            if abs(diff) > 0.01:
                errors.append(f"{month_label}: {label} difference is {diff:,.2f}.")
    return errors


def _field_exists(doctype, fieldname):
    return bool(frappe.db.exists("DocField", {"parent": doctype, "fieldname": fieldname}))


def _resolve_account(account_code, company):
    if frappe.db.exists("Account", account_code):
        return account_code

    filters = {"company": company} if _field_exists("Account", "company") else {}

    if _field_exists("Account", "account_number"):
        name = frappe.db.get_value("Account", {**filters, "account_number": account_code}, "name")
        if name:
            return name

    candidates = frappe.get_all(
        "Account",
        filters=filters,
        fields=["name", "account_name"],
        limit_page_length=1000,
    )
    for row in candidates:
        if cstr(row.name).startswith(account_code):
            return row.name
    for row in candidates:
        if cstr(row.account_name).startswith(account_code):
            return row.name
    return None


def _get_resolved_cost_center_for_row(
    cost_center_code,
    company,
    row_idx,
    warnings,
    errors,
    cache,
    cost_center_was_explicit=False,
):
    """Resolve one Excel cost-center code to an ERPNext Cost Center document name.

    Main business rule:
      - Explicit Excel codes like WON/OTH/KRR must match Cost Center Number.
      - Blank or missing Excel cost center defaults to HO.
      - If a hyphen-ending text token does not exist as a Cost Center Number, it
        is assumed to be part of the description, not a real cost center, and the
        row defaults to HO.
    """
    requested_code = cstr(cost_center_code).strip().upper() or "HO"
    cc_key = (requested_code, company)
    if cc_key in cache:
        resolved = cache[cc_key]
    else:
        resolved = _resolve_cost_center(requested_code, company)
        cache[cc_key] = resolved

    if resolved:
        return resolved

    # The Excel text had a final hyphen-token, but it is not a valid Cost Center
    # Number in ERPNext. Treat it as descriptive text and default the row to HO.
    # Do not hard-code aliases like BANK -> HO or LEASES -> HO; the only default
    # is the formal HO cost center.
    if cost_center_was_explicit and requested_code != "HO":
        ho_key = ("HO", company)
        if ho_key in cache:
            ho = cache[ho_key]
        else:
            ho = _resolve_cost_center("HO", company)
            cache[ho_key] = ho

        if ho:
            warnings.append(
                f"Row {row_idx}: Excel text ended with '{requested_code}', but no active Cost Center "
                f"with Cost Center Number '{requested_code}' exists for company {company}. "
                "The row was defaulted to HO."
            )
            return ho

    errors.append(
        f"Row {row_idx}: Cost Center {requested_code} could not be found for company {company}. "
        "Please create/update a Cost Center with this Cost Center Number, or make sure HO exists."
    )
    return None


def _resolve_cost_center(cost_center_code, company):
    """Resolve by Cost Center Number first. Return the Cost Center document name.

    Your Cost Center master stores short import codes such as WON, OTH and HO in
    the custom/standard field `cost_center_number`, while the document name is
    longer, e.g. `WON - Wonderfontein - ISA`. The child Link field must receive
    the document name.
    """
    cost_center_code = cstr(cost_center_code).strip().upper() or "HO"

    filters = {"company": company} if _field_exists("Cost Center", "company") else {}
    if _field_exists("Cost Center", "disabled"):
        filters["disabled"] = 0

    if _field_exists("Cost Center", "cost_center_number"):
        name = frappe.db.get_value(
            "Cost Center",
            {**filters, "cost_center_number": cost_center_code},
            "name",
        )
        if name:
            return name

    # Exact document-name fallback only. Avoid fuzzy token matching because it can
    # mistake description words for cost centers.
    if frappe.db.exists("Cost Center", cost_center_code):
        cc_company = frappe.db.get_value("Cost Center", cost_center_code, "company") if _field_exists("Cost Center", "company") else company
        disabled = frappe.db.get_value("Cost Center", cost_center_code, "disabled") if _field_exists("Cost Center", "disabled") else 0
        if cc_company == company and not disabled:
            return cost_center_code

    return None

def _build_summary_html(parsed, title="GL Historical Import"):
    controls = parsed.get("excel_controls") or {}
    imported = parsed.get("imported_totals") or {}
    month_labels = sorted(set(controls.keys()) | set(imported.keys()), key=_month_sort_key)

    rows_html = ""
    total_excel_income = total_excel_expense = total_excel_profit = 0
    total_calc_income = total_calc_expense = total_calc_profit = 0

    for month_label in month_labels:
        excel = controls.get(month_label, {})
        calc = imported.get(month_label, {})
        ei = flt(excel.get("income"), 2)
        ee = flt(excel.get("expense"), 2)
        ep = flt(excel.get("profit_loss"), 2)
        ci = flt(calc.get("income"), 2)
        ce = flt(calc.get("expense"), 2)
        cp = flt(calc.get("profit_loss"), 2)
        di = ci - ei
        de = ce - ee
        dp = cp - ep

        total_excel_income += ei
        total_excel_expense += ee
        total_excel_profit += ep
        total_calc_income += ci
        total_calc_expense += ce
        total_calc_profit += cp

        rows_html += f"""
            <tr>
                <td>{frappe.utils.escape_html(month_label)}</td>
                <td class="text-right">{_money(ei)}</td>
                <td class="text-right">{_money(ci)}</td>
                <td class="text-right { _diff_class(di) }">{_money(di)}</td>
                <td class="text-right">{_money(ee)}</td>
                <td class="text-right">{_money(ce)}</td>
                <td class="text-right { _diff_class(de) }">{_money(de)}</td>
                <td class="text-right">{_money(ep)}</td>
                <td class="text-right">{_money(cp)}</td>
                <td class="text-right { _diff_class(dp) }">{_money(dp)}</td>
            </tr>
        """

    ti = total_calc_income - total_excel_income
    te = total_calc_expense - total_excel_expense
    tp = total_calc_profit - total_excel_profit
    rows_html += f"""
        <tr style="font-weight:700; background:#f8f8f8;">
            <td>Total</td>
            <td class="text-right">{_money(total_excel_income)}</td>
            <td class="text-right">{_money(total_calc_income)}</td>
            <td class="text-right { _diff_class(ti) }">{_money(ti)}</td>
            <td class="text-right">{_money(total_excel_expense)}</td>
            <td class="text-right">{_money(total_calc_expense)}</td>
            <td class="text-right { _diff_class(te) }">{_money(te)}</td>
            <td class="text-right">{_money(total_excel_profit)}</td>
            <td class="text-right">{_money(total_calc_profit)}</td>
            <td class="text-right { _diff_class(tp) }">{_money(tp)}</td>
        </tr>
    """

    errors_html = "".join([f"<li>{frappe.utils.escape_html(e)}</li>" for e in parsed.get("errors") or []])
    warnings_html = "".join([f"<li>{frappe.utils.escape_html(w)}</li>" for w in parsed.get("warnings") or []])

    status = parsed.get("status") or "Draft"
    row_count = parsed.get("row_count") or 0

    return f"""
        <div class="glh-summary">
            <h4 style="margin-top:0;">{frappe.utils.escape_html(title)}</h4>
            <p>
                <b>Status:</b> {frappe.utils.escape_html(status)}<br>
                <b>Rows:</b> {row_count}
            </p>
            <table class="table table-bordered table-condensed">
                <thead>
                    <tr>
                        <th>Month</th>
                        <th class="text-right">Excel Income</th>
                        <th class="text-right">Imported Income</th>
                        <th class="text-right">Income Diff</th>
                        <th class="text-right">Excel Expense</th>
                        <th class="text-right">Imported Expense</th>
                        <th class="text-right">Expense Diff</th>
                        <th class="text-right">Excel P/L</th>
                        <th class="text-right">Imported P/L</th>
                        <th class="text-right">P/L Diff</th>
                    </tr>
                </thead>
                <tbody>{rows_html}</tbody>
            </table>
            {f'<div class="alert alert-danger"><b>Errors</b><ul>{errors_html}</ul></div>' if errors_html else ''}
            {f'<div class="alert alert-warning"><b>Warnings</b><ul>{warnings_html}</ul></div>' if warnings_html else ''}
        </div>
    """


def _month_sort_key(label):
    parsed = _parse_month_year(label)
    if parsed:
        return (parsed[1], parsed[0])
    return (9999, 99)


def _money(value):
    return f"{flt(value, 2):,.2f}"


def _diff_class(value):
    return "text-success" if abs(flt(value, 2)) <= 0.01 else "text-danger"
