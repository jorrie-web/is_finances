# -*- coding: utf-8 -*-
"""IS Creditors Recon DocType server logic (Frappe v16).

Goals
- Import Supplier Sage Transactions (Excel) into child table `processed_transactions`.
- Parse Supplier Statement (PDF) and reconcile child rows.

Important constraints
- NO custom API / RPC methods (no frappe.whitelist, no frappe.call endpoints).
- Runs via standard DocType lifecycle (validate/before_save/etc.).
- Always uses the file attached in the field `sage_transactions`.
"""

from __future__ import annotations

# pyright: reportMissingImports=false, reportGeneralTypeIssues=false, reportAttributeAccessIssue=false

import hashlib
import os
import re
import tempfile
import zipfile
from dataclasses import dataclass
from decimal import Decimal, ROUND_HALF_UP
from typing import Dict, List, Set, Tuple

import frappe
from frappe.model.document import Document
from frappe.utils import cint, cstr


# ------------------------
# Helpers
# ------------------------

MONTH_REF_RE = r"(?:JAN|FEB|MAR|APR|MAY|JUN|JUL|AUG|SEP|OCT|NOV|DEC)\d{2,6}"


def _money(v) -> float:
    """Convert common excel/pdf values to float with 2dp rounding."""
    if v is None or v == "":
        return 0.0
    try:
        d = Decimal(str(v)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
        return float(d)
    except Exception:
        s = cstr(v)
        s = s.replace("\u00a0", " ")
        s = s.replace(" ", "")
        s = s.replace(",", "")
        if not s:
            return 0.0
        try:
            d = Decimal(s).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
            return float(d)
        except Exception:
            return 0.0


def _sha256_file(path: str) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def _coerce_statement_date(v) -> str | None:
    s = cstr(v or "").strip()
    if not s:
        return None

    if re.fullmatch(r"\d{4}-\d{2}-\d{2}", s):
        return s

    m = re.fullmatch(r"(\d{2})[./-](\d{2})[./-](\d{4})", s)
    if m:
        dd, mm, yyyy = m.groups()
        return f"{yyyy}-{mm}-{dd}"

    m = re.fullmatch(r"(\d{4})(\d{2})(\d{2})", s)
    if m:
        yyyy, mm, dd = m.groups()
        return f"{yyyy}-{mm}-{dd}"

    try:
        from frappe.utils import getdate
        d = getdate(s)
        return d.isoformat() if d else None
    except Exception:
        return None


def resolve_attached_file_path(file_url: str) -> str:
    """Resolve an Attach field value to an absolute on-disk path."""
    url = cstr(file_url).strip()
    if not url:
        frappe.throw("Please attach a file first.")

    if url.startswith("./"):
        url2 = url[2:]
        first_slash = url2.find("/")
        if first_slash != -1:
            url = url2[first_slash:]
        else:
            url = "/" + url2

    if url.startswith("http://") or url.startswith("https://"):
        from urllib.parse import urlparse
        url = urlparse(url).path

    if url.startswith("/private/files/"):
        rel = url[len("/private/files/"):]
        abs_path = frappe.get_site_path("private", "files", rel)
    elif url.startswith("/files/"):
        rel = url[len("/files/"):]
        abs_path = frappe.get_site_path("public", "files", rel)
    else:
        frappe.throw(f"Invalid attachment path: {file_url}")

    if not os.path.exists(abs_path):
        frappe.throw(f"Attached file not found on server disk. Field value: {file_url}")

    return abs_path


def _sanitize_xlsx_for_openpyxl(src_path: str) -> str:
    """Create a sanitized copy of an XLSX that openpyxl can read."""
    tmpdir = tempfile.mkdtemp(prefix="is_recon_xlsx_")
    extract_dir = os.path.join(tmpdir, "unz")
    os.makedirs(extract_dir, exist_ok=True)

    with zipfile.ZipFile(src_path, "r") as zin:
        zin.extractall(extract_dir)

    replacements = {
        "WindowWidth": "windowWidth",
        "WindowHeight": "windowHeight",
        "WindowX": "windowX",
        "WindowY": "windowY",
        "TabRatio": "tabRatio",
        "FirstSheet": "firstSheet",
        "ActiveTab": "activeTab",
        "firstPageNo": "firstPageNumber",
    }

    def fix_xml(path: str) -> None:
        if not os.path.exists(path):
            return
        try:
            s = open(path, "rb").read().decode("utf-8")
        except Exception:
            return
        for k, v in replacements.items():
            s = s.replace(k + "=", v + "=")
        open(path, "wb").write(s.encode("utf-8"))

    fix_xml(os.path.join(extract_dir, "xl", "workbook.xml"))
    wsdir = os.path.join(extract_dir, "xl", "worksheets")
    if os.path.isdir(wsdir):
        for fn in os.listdir(wsdir):
            if fn.endswith(".xml"):
                fix_xml(os.path.join(wsdir, fn))

    fixed_path = os.path.join(tmpdir, "fixed.xlsx")
    with zipfile.ZipFile(fixed_path, "w", compression=zipfile.ZIP_DEFLATED) as zout:
        for root, _, files in os.walk(extract_dir):
            for f in files:
                p = os.path.join(root, f)
                arc = os.path.relpath(p, extract_dir)
                zout.write(p, arc)

    return fixed_path


@dataclass
class StatementEntry:
    date_str: str
    reference: str
    description: str = ""
    debit: float = 0.0
    credit: float = 0.0
    amount: float = 0.0
    source_row: int = 0
    source_page: int = 0
    match_key: str = ""
    source_text: str = ""


def _parse_statement_pdf(pdf_path: str, parser: dict | None = None) -> List[StatementEntry]:
    parser = parser or {}

    def _extract_text() -> str:
        extractor = cstr(parser.get("extractor") or "auto").strip().lower() or "auto"
        page_start = parser.get("page_start")
        page_end = parser.get("page_end")

        def _slice_pages(pages):
            if not page_start and not page_end:
                return pages
            s = max(int(page_start or 1) - 1, 0)
            e = int(page_end) if page_end else len(pages)
            return pages[s:e]

        if extractor in ("auto", "pypdf"):
            try:
                from pypdf import PdfReader  # type: ignore
                reader = PdfReader(pdf_path)
                pages = _slice_pages(reader.pages)
                txt = "\n".join([(p.extract_text() or "") for p in pages])
                if txt.strip():
                    return txt
                if extractor == "pypdf":
                    return txt
            except Exception:
                if extractor == "pypdf":
                    raise

        if extractor in ("auto", "pdfplumber"):
            try:
                import pdfplumber  # type: ignore
                parts = []
                with pdfplumber.open(pdf_path) as pdf:
                    for p in _slice_pages(pdf.pages):
                        parts.append(p.extract_text() or "")
                txt = "\n".join(parts)
                if txt.strip():
                    return txt
                if extractor == "pdfplumber":
                    return txt
            except Exception:
                if extractor == "pdfplumber":
                    raise

        if extractor in ("auto", "ocr"):
            try:
                import pypdfium2 as pdfium  # type: ignore
                import pytesseract  # type: ignore

                pdf = pdfium.PdfDocument(pdf_path)
                start = max(int(page_start or 1) - 1, 0)
                end = int(page_end) if page_end else len(pdf)
                parts = []
                for i in range(start, min(end, len(pdf))):
                    page = pdf[i]
                    bitmap = page.render(scale=2.0)
                    pil_image = bitmap.to_pil()
                    parts.append(pytesseract.image_to_string(pil_image))
                txt = "\n".join(parts)
                if txt.strip():
                    return txt
            except Exception as e:
                if extractor == "ocr":
                    frappe.throw(f"Could not OCR PDF statement: {e}")

        frappe.throw("Could not read PDF statement. No usable text layer was found and OCR did not return text.")
        return ""

    def _compile_or_throw(pattern: str, fieldname: str):
        try:
            return re.compile(pattern)
        except re.error as e:
            parser_name = cstr(parser.get("parser_name") or parser.get("name") or "Unknown Parser")
            frappe.throw(
                f"Invalid {fieldname} in Statement Parser '{parser_name}': {e}\n\n{pattern}",
                title="Statement Parser Regex Error",
            )

    def _parse_amount_str(s: str) -> float:
        s = cstr(s or "").strip()
        if not s:
            return 0.0
        neg = False
        s = re.sub(r"^[\)\]\}]+|[\(\[\{]+$", "", s).strip()
        if s.startswith("(") and s.endswith(")"):
            neg = True
            s = s[1:-1].strip()
        s = re.sub(r"(?<!\d)[\)\]\}]+|[\(\[\{]+(?!\d)", "", s)
        s = s.replace("\u00a0", " ").replace(" ", "").replace(",", "")
        if s.startswith("-"):
            neg = True
        val = _money(s)
        return -abs(val) if neg else val

    def _extract_reference_from_line(line: str, regex_ref: str = "") -> str:
        line = cstr(line or "")
        regex_ref = cstr(regex_ref or "").strip().upper()

        if re.fullmatch(r"\d{8,9}", regex_ref):
            return regex_ref
        if re.fullmatch(r"DISC", regex_ref):
            return regex_ref
        if re.fullmatch(MONTH_REF_RE, regex_ref):
            return regex_ref

        m = re.search(r"\b(\d{8,9})\b(?=\s+ZAR\b)", line, re.IGNORECASE)
        if m:
            return m.group(1)

        m = re.search(rf"\b(DISC|{MONTH_REF_RE})\b", line, re.IGNORECASE)
        if m:
            return m.group(1).upper()

        if re.fullmatch(r"\d{6}[A-Za-z]{2}", regex_ref):
            regex_ref = regex_ref[:6]
        if re.fullmatch(r"\d{6}", regex_ref):
            regex_choice = regex_ref
        else:
            regex_choice = ""

        tokens = re.split(r"\s+", line.strip())
        for i, tok in enumerate(tokens):
            tok_u = cstr(tok).upper()
            if re.fullmatch(r"\d{6}[A-Za-z]{2}", tok_u):
                return tok_u[:6]
            if re.fullmatch(r"\d{6}", tok_u):
                prev_tok = tokens[i - 1] if i > 0 else ""
                next_tok = tokens[i + 1] if i + 1 < len(tokens) else ""
                if re.fullmatch(r"[A-Za-z]{2}", prev_tok) or re.fullmatch(r"[A-Za-z]{2}", next_tok):
                    return tok_u

        if regex_choice:
            return regex_choice

        first_six = re.search(r"\b(\d{6})\b", line)
        if first_six:
            return first_six.group(1)

        merged = re.search(r"\b(\d{6})([A-Za-z]{2})\b", line)
        if merged:
            return merged.group(1)

        return regex_ref

    def _sanitize_statement_line(line: str) -> str:
        line = cstr(line or "")
        if not line:
            return ""
        line = line.replace("\u00a0", " ")
        line = line.replace("«", " ").replace("»", " ")
        line = re.sub(r"\s*\|\s*(?=ZAR\b)", " ", line, flags=re.IGNORECASE)
        line = re.sub(r"(?<=\d)\s*\|\s*(?=\d{1,3}(?:[,\s]\d{3})*\.\d{2}\b)", " ", line)
        line = re.sub(r"(?<=\d\.\d{2})[\|\)\]\}:;.,]+(?=\s|$)", "", line)
        line = re.sub(r"[\(\[\{]+(?=\d{1,3}(?:[\s,]\d{3})*\.\d{2})", "", line)
        line = re.sub(r"\s+", " ", line).strip()
        return line

    def _extract_dr_cr_from_line(line: str) -> Tuple[float | None, float | None]:
        clean_line = _sanitize_statement_line(line)
        vals = re.findall(r"(?<!\d)-?\d{1,3}(?:[ ,]\d{3})*\.\d{2}(?!\d)", clean_line)
        if len(vals) < 2:
            return None, None
        dr = _parse_amount_str(vals[-2])
        cr = _parse_amount_str(vals[-1])
        return dr, cr

    def _extract_amount_from_line(line: str, regex_amount: str = "") -> str:
        regex_amount = cstr(regex_amount or "").strip()
        clean_line = _sanitize_statement_line(line)
        money_tokens = re.findall(r"(?<!\d)-?\d{1,3}(?:[\s,]\d{3})*\.\d{2}(?!\d)", clean_line)
        if len(money_tokens) >= 2:
            return money_tokens[-2]
        if money_tokens:
            return money_tokens[0]
        return regex_amount

    full_text = _extract_text()
    txt = (full_text or "").replace("\u00a0", " ")

    collapse = cint(parser.get("collapse_whitespace") or 0) == 1
    if collapse or not parser:
        lines = [re.sub(r"\s+", " ", ln).strip() for ln in txt.splitlines()]
    else:
        lines = [ln.rstrip("\n") for ln in txt.splitlines()]

    skip_pat = _compile_or_throw(parser.get("skip_lines_regex"), "skip_lines_regex") if parser.get("skip_lines_regex") else None
    stop_pat = _compile_or_throw(parser.get("stop_at_regex"), "stop_at_regex") if parser.get("stop_at_regex") else None

    row_re = cstr(parser.get("row_regex") or "").strip()
    sign_rule = cstr(parser.get("amount_sign_rule") or "").strip().lower()

    if row_re:
        pat = _compile_or_throw(row_re, "row_regex")
    else:
        date_re = r"(?P<date>\d{2}[A-Za-z]{3}\d{2})"
        money_re = r"-?\d{1,3}(?:[\s,]\d{3})*\.\d{2}"
        pat = re.compile(
            rf"^{date_re}\s+(?P<ref>\d{{6}})(?P<doc_type>[A-Za-z]{{2}}|\s+[A-Za-z]{{2}})?\s+.*?(?P<m1>{money_re})\s+(?P<m2>{money_re})$"
        )

    out: List[StatementEntry] = []
    payment_fallback = re.compile(
        rf"^(?P<date>\d{{2}}\.\d{{2}}\.\d{{4}})\s+\|?\s*.*?(?P<ref>DISC|{MONTH_REF_RE})\b.*?(?:\bZAR\b\s+)?(?P<m1>\d{{1,3}}(?:[\s,]\d{{3}})*\.\d{{2}}|\d+\.\d{{2}})[\|\)\]\}}:;.,]*\s+(?P<m2>\d{{1,3}}(?:[\s,]\d{{3}})*\.\d{{2}}|\d+\.\d{{2}})[\|\)\]\}}:;.,]*\s*$",
        re.IGNORECASE,
    )
    invoice_fallback = re.compile(
        rf"^(?P<date>\d{{2}}\.\d{{2}}\.\d{{4}})\s+\|?\s*.*?(?P<ref>\d{{8,9}}|DISC|{MONTH_REF_RE})\b.*?(?:\bZAR\b\s+)?(?P<m1>\d{{1,3}}(?:[\s,]\d{{3}})*\.\d{{2}}|\d+\.\d{{2}})[\|\)\]\}}:;.,]*\s+(?P<m2>\d{{1,3}}(?:[\s,]\d{{3}})*\.\d{{2}}|\d+\.\d{{2}})[\|\)\]\}}:;.,]*\s*$",
        re.IGNORECASE,
    )

    for ln in lines:
        if not ln:
            continue
        clean_ln = _sanitize_statement_line(ln)
        if skip_pat and (skip_pat.search(ln) or skip_pat.search(clean_ln)):
            continue
        if stop_pat and (stop_pat.search(ln) or stop_pat.search(clean_ln)):
            break

        m = pat.match(ln) or pat.match(clean_ln)
        if not m:
            m = payment_fallback.match(clean_ln) or invoice_fallback.match(clean_ln)
            if not m:
                continue

        gd = m.groupdict()
        date_str = cstr(gd.get("date") or "").strip()
        ref = _extract_reference_from_line(clean_ln or ln, gd.get("ref") or "")
        if not ref:
            continue

        amount_style = cstr(parser.get("amount_style") or "").strip().lower()
        if amount_style == "pick_non_zero_last_two":
            dr, cr = _extract_dr_cr_from_line(ln)
            if dr is None or cr is None:
                continue
            if abs(dr) > 0 and abs(cr) > 0 and _close(abs(dr), abs(cr), 0.02):
                continue
            if abs(dr) > 0 and _close(abs(cr), 0.0, 0.02):
                amt = abs(dr)
            elif abs(cr) > 0 and _close(abs(dr), 0.0, 0.02):
                amt = abs(cr)
            else:
                continue
        else:
            amt_raw = _extract_amount_from_line(ln, gd.get("amount") or gd.get("m1") or "")
            if not amt_raw:
                continue
            amt = _parse_amount_str(amt_raw)

            if sign_rule == "credit_positive":
                amt = abs(amt)
            elif sign_rule == "debit_positive" and amt < 0:
                amt = abs(amt)
            elif sign_rule == "absolute_non_zero":
                amt = abs(amt)

        debit = abs(amt) if amt > 0 else 0.0
        credit = abs(amt) if amt < 0 else 0.0
        out.append(StatementEntry(
            date_str=date_str,
            reference=ref,
            description=clean_ln or ln,
            debit=debit,
            credit=credit,
            amount=amt,
            source_row=len(out) + 1,
            source_page=0,
            match_key=f"{_normalize_ref(ref)}|{abs(amt):.2f}" if ref else "",
            source_text=clean_ln or ln,
        ))

    return out


def _resolve_statement_parser(doc: "ISCreditorsRecon") -> dict | None:
    parser_name = cstr(getattr(doc, "statement_parser", "")).strip()
    if parser_name:
        parser_doc = frappe.db.get_value("IS Supplier Statement Parser", parser_name, "*", as_dict=True)
        if parser_doc and cint(parser_doc.get("is_active") or 0) == 1:
            return parser_doc

    supplier_code = cstr(getattr(doc, "supplier_code", "")).strip()
    if supplier_code:
        rows = frappe.get_all(
            "IS Supplier Statement Parser",
            filters={"is_active": 1, "match_on": "Supplier Code", "supplier_code": supplier_code},
            fields=["*"],
            order_by="priority desc, modified desc",
            limit=1,
        )
        if rows:
            return rows[0]

    supplier = cstr(getattr(doc, "supplier", "")).strip()
    if supplier:
        rows = frappe.get_all(
            "IS Supplier Statement Parser",
            filters={"is_active": 1, "match_on": "Supplier", "supplier": supplier},
            fields=["*"],
            order_by="priority desc, modified desc",
            limit=1,
        )
        if rows:
            return rows[0]

    return None


def _normalize_ref(ref: str) -> str:
    return re.sub(r"\s+", "", cstr(ref or "")).upper()


def _close(a: float, b: float, tol: float = 0.01) -> bool:
    da = Decimal(str(a or 0.0))
    db = Decimal(str(b or 0.0))
    dt = Decimal(str(tol))
    return abs(da - db) <= dt


# ------------------------
# DocType
# ------------------------

class ISCreditorsRecon(Document):
    """DocType: IS Creditors Recon"""

    def validate(self):
        self._import_excel_if_needed()
        self._reconcile_against_pdf_if_present()

    def _import_excel_if_needed(self) -> None:
        if not self.sage_transactions:
            return

        excel_path = resolve_attached_file_path(self.sage_transactions)
        file_hash = _sha256_file(excel_path)

        if self.import_hash and self.import_hash == file_hash and self.processed_transactions:
            return

        try:
            import openpyxl  # type: ignore
        except Exception as e:
            frappe.throw(f"openpyxl is required to import Excel files: {e}")

        fixed_path = _sanitize_xlsx_for_openpyxl(excel_path)

        try:
            wb = openpyxl.load_workbook(fixed_path, data_only=True)
            ws = wb.active
        except Exception as e:
            frappe.throw(f"Could not read Excel file: {e}")

        header_row = None
        for r in range(1, min(ws.max_row, 80) + 1):
            v1 = cstr(ws.cell(r, 1).value).strip().lower()
            v3 = cstr(ws.cell(r, 3).value).strip().lower()
            if v1 == "supplier" and v3 == "date":
                header_row = r
                break

        if not header_row:
            frappe.throw("Could not locate header row in the Excel file (expected 'Supplier' and 'Date' columns).")

        from_date = ws.cell(5, 2).value
        to_date = ws.cell(6, 2).value
        if from_date:
            self.opening_date = from_date
        if to_date:
            self.closing_date = to_date

        self.set("processed_transactions", [])

        opening_balance_set = False
        closing_balance_set = False

        for r in range(header_row + 1, ws.max_row + 1):
            supplier = ws.cell(r, 1).value
            date_val = ws.cell(r, 3).value
            ref_val = ws.cell(r, 5).value
            code_val = ws.cell(r, 7).value
            debit_val = ws.cell(r, 10).value
            credit_val = ws.cell(r, 13).value
            desc_val = ws.cell(r, 14).value

            if isinstance(ref_val, str) and ref_val.strip().lower() == "opening balance":
                bal = _money(credit_val)
                self.o_balance = bal
                opening_balance_set = True
                self.append("processed_transactions", {
                    "tran_date": self.opening_date or date_val,
                    "reference": "OPENING",
                    "tr_code": "BAL",
                    "description": "Opening Balance",
                    "debit": 0,
                    "credit": bal,
                    "amount": -abs(bal),
                    "source_row": r,
                    "reconciled": 1,
                    "is_opening_balance": 1,
                    "is_closing_balance": 0,
                    "match_key": "OPENING",
                    "match_key_tracked": "OPENING",
                    "match_key_issue_type": "",
                    "match_key_issue_notes": "",
                    "unmatched_category": "",
                    "unmatched_notes": "",
                    "manual_matched": 0,
                })
                continue

            if isinstance(ref_val, str) and ref_val.strip().lower() == "closing balance":
                bal = _money(credit_val)
                self.closing_balance = bal
                closing_balance_set = True
                self.append("processed_transactions", {
                    "tran_date": self.closing_date or date_val,
                    "reference": "CLOSING",
                    "tr_code": "BAL",
                    "description": "Closing Balance",
                    "debit": 0,
                    "credit": bal,
                    "amount": -abs(bal),
                    "source_row": r,
                    "reconciled": 1,
                    "is_opening_balance": 0,
                    "is_closing_balance": 1,
                    "match_key": "CLOSING",
                    "match_key_tracked": "CLOSING",
                    "match_key_issue_type": "",
                    "match_key_issue_notes": "",
                    "unmatched_category": "",
                    "unmatched_notes": "",
                    "manual_matched": 0,
                })
                continue

            if not supplier and not date_val and not ref_val and not desc_val:
                continue
            if isinstance(supplier, str) and supplier.lower().startswith("supplier:"):
                continue

            supplier_s = cstr(supplier).strip().lower()
            desc_s = cstr(desc_val).strip().lower()
            if supplier_s.startswith("sage 200 evolution"):
                continue
            if not date_val and not ref_val and not desc_s and supplier_s:
                continue

            debit = _money(debit_val)
            credit = _money(credit_val)
            amount = debit - credit

            ref = cstr(ref_val).strip()
            if not ref:
                ref = f"ROW-{r}"

            audit_no = ""
            if isinstance(desc_val, str):
                m = re.search(r"\b(AD\d{3,}|TDQ\b|[A-Z]{2,}\d{3,})\b", desc_val.upper())
                if m:
                    audit_no = m.group(1)

            match_key = f"{_normalize_ref(ref)}|{abs(_money(credit) or abs(amount)):.2f}"

            self.append("processed_transactions", {
                "tran_date": date_val,
                "reference": ref,
                "tr_code": cstr(code_val).strip(),
                "description": cstr(desc_val).strip(),
                "debit": debit,
                "credit": credit,
                "amount": amount,
                "audit_no": audit_no,
                "match_key": match_key,
                "match_key_tracked": match_key,
                "match_key_issue_type": "",
                "match_key_issue_notes": "",
                "unmatched_category": "",
                "unmatched_notes": "",
                "source_row": r,
                "reconciled": 0,
                "is_opening_balance": 0,
                "is_closing_balance": 0,
                "manual_matched": 0,
            })

        if not opening_balance_set and self.o_balance is None:
            self.o_balance = 0
        if not closing_balance_set and self.closing_balance is None:
            self.closing_balance = 0

        self.import_hash = file_hash
        self.imported_on = frappe.utils.now_datetime()

    def _populate_statement_transactions(self, entries: List[StatementEntry], parser: dict | None = None) -> None:
        prev_rows_by_key: Dict[str, Dict[str, object]] = {}
        for row in list(getattr(self, "statement_transactions", []) or []):
            base_key = cstr(getattr(row, "match_key", "")).strip()
            tracked_key = cstr(getattr(row, "match_key_tracked", "")).strip()
            prev_key = tracked_key or base_key
            if prev_key:
                prev_rows_by_key[prev_key] = {
                    "manual_matched": cint(getattr(row, "manual_matched", 0)),
                    "reconciled": cint(getattr(row, "reconciled", 0)),
                    "match_key_tracked": tracked_key,
                    "unmatched_category": cstr(getattr(row, "unmatched_category", "")),
                    "unmatched_notes": cstr(getattr(row, "unmatched_notes", "")),
                    "match_key_issue_type": cstr(getattr(row, "match_key_issue_type", "")),
                    "match_key_issue_notes": cstr(getattr(row, "match_key_issue_notes", "")),
                }

        self.set("statement_transactions", [])
        parser_used = cstr((parser or {}).get("parser_name") or (parser or {}).get("name") or "DEFAULT")
        for entry in entries:
            amount = _money(entry.amount)
            debit = _money(entry.debit) if entry.debit else (_money(abs(amount)) if amount > 0 else 0.0)
            credit = _money(entry.credit) if entry.credit else (_money(abs(amount)) if amount < 0 else 0.0)
            match_key = entry.match_key or (f"{_normalize_ref(entry.reference)}|{abs(amount):.2f}" if entry.reference else "")
            prev_row = prev_rows_by_key.get(match_key, {})
            was_manual = cint(prev_row.get("manual_matched") or 0) == 1
            self.append("statement_transactions", {
                "tran_date": _coerce_statement_date(entry.date_str),
                "reference": entry.reference,
                "description": entry.description or entry.source_text or "",
                "debit": debit,
                "credit": credit,
                "amount": amount,
                "source_row": cint(entry.source_row),
                "source_page": cint(entry.source_page),
                "match_key": match_key,
                "match_key_tracked": cstr(prev_row.get("match_key_tracked") or match_key),
                "reconciled": 1 if was_manual else 0,
                "manual_matched": 1 if was_manual else 0,
                "is_opening_balance": 0,
                "is_closing_balance": 0,
                "unmatched_category": cstr(prev_row.get("unmatched_category") or ""),
                "unmatched_notes": cstr(prev_row.get("unmatched_notes") or ""),
                "match_key_issue_type": cstr(prev_row.get("match_key_issue_type") or ""),
                "match_key_issue_notes": cstr(prev_row.get("match_key_issue_notes") or ""),
                "parser_name": parser_used,
                "source_text": entry.source_text or entry.description or "",
            })

    def _compute_recon_summary_totals(self) -> None:
        def row_total(rows, matched: bool) -> float:
            total = 0.0
            for row in rows:
                if bool(cint(getattr(row, "reconciled", 0))) != matched:
                    continue
                if cint(getattr(row, "is_opening_balance", 0)) or cint(getattr(row, "is_closing_balance", 0)):
                    continue
                debit = _money(getattr(row, "debit", 0))
                credit = _money(getattr(row, "credit", 0))
                amount = _money(getattr(row, "amount", 0))
                total += abs(credit) if credit else (abs(debit) if debit else abs(amount))
            return round(total, 2)

        self.erp_matched_total = row_total(list(getattr(self, "processed_transactions", []) or []), True)
        self.erp_unmatched_total = row_total(list(getattr(self, "processed_transactions", []) or []), False)
        self.statement_matched_total = row_total(list(getattr(self, "statement_transactions", []) or []), True)
        self.statement_unmatched_total = row_total(list(getattr(self, "statement_transactions", []) or []), False)
        self.closing_variance = round(_money(getattr(self, "closing_balance", 0)) - _money(getattr(self, "statement_closing_balance", 0)), 2)

    def _reconcile_against_pdf_if_present(self) -> None:
        if not self.sup_statement:
            self._compute_recon_summary_totals()
            return
        if not self.processed_transactions:
            self._compute_recon_summary_totals()
            return

        pdf_path = resolve_attached_file_path(self.sup_statement)
        parser = _resolve_statement_parser(self)
        entries = _parse_statement_pdf(pdf_path, parser=parser)
        if hasattr(self, "parser_used"):
            self.parser_used = cstr((parser or {}).get("parser_name") or (parser or {}).get("name") or "DEFAULT")

        if entries:
            self._populate_statement_transactions(entries, parser=parser)
        elif not getattr(self, "statement_transactions", None):
            self.set("statement_transactions", [])
            self.statement_opening_balance = 0
            self.statement_closing_balance = 0
            self._compute_recon_summary_totals()
            return

        tol = 0.10 if cstr((parser or {}).get("extractor") or "").strip().lower() == "ocr" else 0.02
        self.statement_opening_balance = _money(getattr(self, "statement_opening_balance", 0))
        self.statement_closing_balance = _money(getattr(self, "statement_closing_balance", 0))

        by_ref: Dict[str, List[Document]] = {}
        by_key: Dict[str, List[Document]] = {}
        for srow in list(getattr(self, "statement_transactions", []) or []):
            srow.reconciled = 1 if cint(getattr(srow, "manual_matched", 0)) else 0
            ref = _normalize_ref(getattr(srow, "reference", ""))
            if ref:
                by_ref.setdefault(ref, []).append(srow)
            key = cstr(getattr(srow, "match_key_tracked", "") or getattr(srow, "match_key", "")).strip()
            if key:
                by_key.setdefault(key, []).append(srow)
            if cint(getattr(srow, "is_opening_balance", 0)):
                self.statement_opening_balance = _money(getattr(srow, "credit", 0) or getattr(srow, "debit", 0) or abs(_money(getattr(srow, "amount", 0))))
                srow.reconciled = 1
            if cint(getattr(srow, "is_closing_balance", 0)):
                self.statement_closing_balance = _money(getattr(srow, "credit", 0) or getattr(srow, "debit", 0) or abs(_money(getattr(srow, "amount", 0))))
                srow.reconciled = 1

        matched_statement_names: Set[str] = {
            cstr(getattr(srow, "name", ""))
            for srow in list(getattr(self, "statement_transactions", []) or [])
            if cint(getattr(srow, "reconciled", 0)) == 1 and cstr(getattr(srow, "name", ""))
        }

        def _srow_id(srow: Document) -> str:
            return cstr(getattr(srow, "name", "")) or f"row-{cint(getattr(srow, 'idx', 0))}"

        def mark_statement_match(candidates: List[Document], target_amt: float) -> bool:
            for srow in candidates:
                sid = _srow_id(srow)
                if sid in matched_statement_names:
                    continue
                s_amt = abs(_money(getattr(srow, "credit", 0)))
                if not s_amt:
                    s_amt = abs(_money(getattr(srow, "debit", 0)))
                if not s_amt:
                    s_amt = abs(_money(getattr(srow, "amount", 0)))
                if _close(s_amt, abs(target_amt), tol):
                    srow.reconciled = 1
                    matched_statement_names.add(sid)
                    return True
            return False

        for row in self.processed_transactions:
            if getattr(row, "is_opening_balance", 0):
                row.reconciled = 1
                continue
            if getattr(row, "is_closing_balance", 0):
                row.reconciled = 1
                continue

            if cint(getattr(row, "manual_matched", 0)) == 1:
                row.reconciled = 1
                row_key = cstr(getattr(row, "match_key_tracked", "") or getattr(row, "match_key", "")).strip()
                if row_key:
                    for srow in by_key.get(row_key, []) or []:
                        sid = _srow_id(srow)
                        if sid in matched_statement_names:
                            continue
                        srow.manual_matched = 1
                        srow.reconciled = 1
                        matched_statement_names.add(sid)
                        break
                if not cstr(getattr(row, "match_key_tracked", "")).strip():
                    row.match_key_tracked = "manually matched"
                continue

            ref = _normalize_ref(getattr(row, "reference", ""))
            debit = _money(getattr(row, "debit", 0))
            credit = _money(getattr(row, "credit", 0))
            target_amt = abs(credit) if credit else (abs(debit) if debit else abs(_money(getattr(row, "amount", 0))))
            matched = False

            row_key = cstr(getattr(row, "match_key_tracked", "") or getattr(row, "match_key", "")).strip()
            if row_key and by_key.get(row_key):
                matched = mark_statement_match(by_key.get(row_key) or [], target_amt)

            if not matched and ref and by_ref.get(ref):
                matched = mark_statement_match(by_ref.get(ref) or [], target_amt)

            if not matched and target_amt:
                unique_amt = []
                for srow in list(getattr(self, "statement_transactions", []) or []):
                    sid = _srow_id(srow)
                    if sid in matched_statement_names:
                        continue
                    s_amt = abs(_money(getattr(srow, "credit", 0)))
                    if not s_amt:
                        s_amt = abs(_money(getattr(srow, "debit", 0)))
                    if not s_amt:
                        s_amt = abs(_money(getattr(srow, "amount", 0)))
                    if _close(s_amt, target_amt, tol):
                        unique_amt.append(srow)
                if len(unique_amt) == 1:
                    sid = _srow_id(unique_amt[0])
                    unique_amt[0].reconciled = 1
                    matched_statement_names.add(sid)
                    matched = True

            row.reconciled = 1 if matched else 0

        if not _money(getattr(self, "statement_closing_balance", 0)):
            stmt_total = 0.0
            for srow in list(getattr(self, "statement_transactions", []) or []):
                if cint(getattr(srow, "is_opening_balance", 0)) or cint(getattr(srow, "is_closing_balance", 0)):
                    continue
                stmt_total += _money(getattr(srow, "debit", 0)) - _money(getattr(srow, "credit", 0))
            self.statement_closing_balance = round(_money(getattr(self, "statement_opening_balance", 0)) + stmt_total, 2)

        self._compute_recon_summary_totals()
